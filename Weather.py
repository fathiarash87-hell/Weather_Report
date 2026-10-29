import requests
import pandas as pd

url = "https://api.open-meteo.com/v1/forecast"

cities = {
    "Tehran": (35.69, 51.39),
    "Shiraz": (29.59, 52.58),
    "Isfahan": (32.65, 51.67),
}

def get_forecast(city, lat, lon):
    params = {
        "latitude": lat,
        "longitude": lon,
        "daily": "temperature_2m_max,temperature_2m_min,precipitation_sum",
        "timezone": "Asia/Tehran",
        "forecast_days": 7,
    }
    response = requests.get(url, params=params, timeout=10)
    response.raise_for_status()
    daily = response.json()["daily"]

    df = pd.DataFrame({
        "Date": daily["time"],
        "Max": daily["temperature_2m_max"],
        "Min": daily["temperature_2m_min"],
        "Rain": daily["precipitation_sum"],
    })
    df["City"] = city
    return df

frames = []

for city, (lat, lon) in cities.items():
    try:
        frames.append(get_forecast(city, lat, lon))
        print(city, "OK")
    except requests.exceptions.RequestException as e:
        print(city, "خطا:", e)

if not frames:
    print("هیچ داده‌ای گرفته نشد، برنامه متوقف شد")
    raise SystemExit

all_data = pd.concat(frames, ignore_index=True)
print(all_data)
print(all_data.shape)

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference


def format_excel(wb):
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        for column in ws.columns:
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in column)
            ws.column_dimensions[column[0].column_letter].width = max_len + 4
        ws.freeze_panes = "A2"


summary = (
    all_data.groupby("City")[["Max", "Min", "Rain"]]
    .agg({"Max": "max", "Min": "min", "Rain": "sum"})
    .round(1)
    .reset_index()
)

# 1) هر شهر یک شیت + یک شیت خلاصه
with pd.ExcelWriter("weather.xlsx") as writer:
    for city in cities:
        city_df = all_data[all_data["City"] == city]
        city_df.to_excel(writer, sheet_name=city, index=False)
    summary.to_excel(writer, sheet_name="Summary", index=False)

# 2) فرمت‌بندی
wb = load_workbook("weather.xlsx")
format_excel(wb)

# 3) نمودار برای هر شهر
from openpyxl.chart.layout import Layout, ManualLayout

for city in cities:
    ws = wb[city]
    last_row = ws.max_row

    chart = LineChart()
    chart.title = f"{city} - 7 day forecast"
    chart.y_axis.title = "Temp (°C)"
    chart.height = 10
    chart.width = 20

    data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    # لجند سمت راست، بیرون از ناحیه‌ی رسم
    chart.legend.position = "r"
    chart.legend.overlay = False

    # ناحیه‌ی رسم رو کوچیک‌تر کن تا برای لجند و عنوان جا بمونه
    chart.plot_area.layout = Layout(
        manualLayout=ManualLayout(
            layoutTarget="inner",
            xMode="edge", yMode="edge",
            x=0.08, y=0.15, w=0.72, h=0.68,
        )
    )

    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.number_format = "mm-dd"
    chart.x_axis.tickLblPos = "low"

    chart.series[0].graphicalProperties.line.solidFill = "C00000"
    chart.series[1].graphicalProperties.line.solidFill = "2E75B6"
    for s in chart.series:
        s.graphicalProperties.line.width = 28000
        s.smooth = False

    ws.add_chart(chart, "G2")

wb.save("weather.xlsx")
print("فایل weather.xlsx ساخته شد")